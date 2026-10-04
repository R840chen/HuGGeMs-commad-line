# -*- coding: utf-8 -*-
"""生成补充表：marker 级别的身份/位置明细 + 汇总页"""
import sys, os, json, collections, re
import numpy as np
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from mapper import read_fasta

# Input FASTA files live in $MARKER_DATA_DIR (default: ./data); outputs go to $MARKER_OUT_DIR (default: this directory)
BASE = os.environ.get('MARKER_DATA_DIR', os.path.join(HERE, 'data'))
OUTDIR = os.environ.get('MARKER_OUT_DIR', HERE)
OUT = os.path.join(OUTDIR, "Supplementary Table S21 - marker gene positional distribution.xlsx")
recs = json.load(open(os.path.join(OUTDIR, 'marker_records.json')))
pos_s = json.load(open(os.path.join(OUTDIR, 'pos_summary.json')))
oric_b = json.load(open(os.path.join(OUTDIR, 'oric_bias.json')))

MARKERS = {'Waltera_2981769': 'Brotolimicola_acetigignens_species_2981769.ffn',
           'Dysosmobacter_2763041': 'Dysosmobacter_hominis_species_2763041.ffn'}
SN = {'W1_GCA_025567015': 'GCA_025567015.1', 'W2_GCA_020687025': 'GCA_020687025.1',
      'D1_GCA_014297375': 'GCA_014297375.1'}
SN2 = {'W1_GCA_025567015': 'W1', 'W2_GCA_020687025': 'W2', 'D1_GCA_014297375': 'D1'}

seqs = {}
for mk, mf in MARKERS.items():
    seqs[mk] = {i: s for i, (_, s) in enumerate(read_fasta(os.path.join(BASE, mf)))}


def gcpct(s):
    s = s.upper(); return round((s.count('G') + s.count('C')) / len(s) * 100, 1)


def parse_ann(h):
    body = h.split(' ', 1)
    if len(body) < 2:
        return '', ''
    tag = body[0].rsplit('_', 1)[-1]
    ann = body[1].split('|')[0].strip()
    return tag, ann


def oric_half(o):
    """oriC 侧一半的 marker 百分比（由十分位分布前 5 个箱求和）"""
    h = o.get('hist')
    if not h:
        return float('nan')
    return sum(h[:5]) / sum(h) * 100


wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Marker positional distribution"
hdr = ['SGB', 'Marker_ID', 'Source_genome', 'Locus_tag', 'Annotation',
       'Conserved_hypothetical', 'Marker_length_bp', 'Marker_GC_pct']
GN = ['GCA_025567015.1', 'GCA_020687025.1', 'GCA_014297375.1']
for g in GN:
    hdr += [f'{g}_identity', f'{g}_coverage', f'{g}_contig', f'{g}_start',
            f'{g}_strand', f'{g}_n_highident_loci']
ws.append(hdr)

for r in sorted(recs, key=lambda x: (x['species'], x['index'])):
    tag, ann = parse_ann(r['header'])
    sq = seqs[r['species']][r['index']]
    row = [r['species'], r['index'] + 1, r['source'], tag, ann,
           'yes' if r['hypothetical'] else 'no', r['len'], gcpct(sq)]
    for gk, g in zip(['W1_GCA_025567015', 'W2_GCA_020687025', 'D1_GCA_014297375'], GN):
        row += [r.get(f'{gk}_id'), r.get(f'{gk}_cov'), r.get(f'{gk}_contig'),
                r.get(f'{gk}_pos'), r.get(f'{gk}_strand'), r.get(f'{gk}_nloc')]
    ws.append(row)

# 汇总页
ws2 = wb.create_sheet("Summary")
rows = [
    ['HuGGeMs marker gene positional distribution — summary', ''],
    ['', ''],
    ['SGB', 'Waltera acetigignens (SGB_2981769)', 'Dysosmobacter hominis (SGB_2763041)'],
    ['Markers analysed', 200, 200],
    ['Marker length, median (bp)', 1935, 1606],
    ['Conserved hypothetical proteins (%)', 30.5, 37.0],
    ['Marker GC content, median (%)', 49.4, 57.7],
    ['Genome GC content (%)', '46.0 / 46.3 (two strains)', 55.7],
    ['', ''],
    ['Genomes analysed', 'GCA_025567015.1; GCA_020687025.1', 'GCA_014297375.1'],
    ['Markers assigned to a genome position', f"{pos_s['W1']['n']} / {pos_s['W2']['n']}", pos_s['D1']['n']],
    ['Contigs carrying markers', f"{pos_s['W1']['ncontig']} / {pos_s['W2']['ncontig']}", pos_s['D1']['ncontig']],
    ['Independent loci (gap > 5 kb)', f"{pos_s['W1']['loci']} / {pos_s['W2']['loci']}", pos_s['D1']['loci']],
    ['Median spacing between adjacent markers (bp)', f"{pos_s['W1']['gaps_median']} / {pos_s['W2']['gaps_median']}", pos_s['D1']['gaps_median']],
    ['Max. markers in any 10-kb window', '4 (2.2%) / 4 (2.2%)', '4 (2.1%)'],
    ['Max. markers in any 20-kb window', '6 (3.4%) / 6 (3.4%)', '6 (3.1%)'],
    ['Max. markers in any 50-kb window', '9 (5.0%) / 8 (4.5%)', '10 (5.1%)'],
    ['Uniformity across contigs (chi-square p)', f"{pos_s['W1']['chi2_p']:.3f} / {pos_s['W2']['chi2_p']:.3f}", f"{pos_s['D1']['chi2_p']:.3f}"],
    ['Uniformity within contigs (KS p)', f"{pos_s['W1']['ks_p']:.3f} / {pos_s['W2']['ks_p']:.3f}", f"{pos_s['D1']['ks_p']:.3f}"],
    ['Uniformity along oriC–ter axis (KS p)', f"{oric_b['W1']['ks_p']:.3f} / {oric_b['W2']['ks_p']:.3f}", f"{oric_b['D1']['ks_p']:.3f}"],
    ['Markers on oriC-proximal half (%)', f"{oric_half(oric_b['W1']):.0f}% / {oric_half(oric_b['W2']):.0f}%", f"{oric_half(oric_b['D1']):.0f}%"],
    ['Estimated abundance bias, oriC:ter = 2:1', f"{oric_b['W1']['bias']['2.0']*100-100:+.2f}% / {oric_b['W2']['bias']['2.0']*100-100:+.2f}%", f"{oric_b['D1']['bias']['2.0']*100-100:+.2f}%"],
    ['Estimated abundance bias, oriC:ter = 4:1', f"{oric_b['W1']['bias']['4.0']*100-100:+.2f}% / {oric_b['W2']['bias']['4.0']*100-100:+.2f}%", f"{oric_b['D1']['bias']['4.0']*100-100:+.2f}%"],
    ['Estimated abundance bias, oriC:ter = 8:1', f"{oric_b['W1']['bias']['8.0']*100-100:+.2f}% / {oric_b['W2']['bias']['8.0']*100-100:+.2f}%", f"{oric_b['D1']['bias']['8.0']*100-100:+.2f}%"],
    ['Multi-copy markers (>=2 loci >= 90% id)', f"{pos_s['W1']['multi']} / {pos_s['W2']['multi']}", pos_s['D1']['multi']],
    ['', ''],
    ['Strain-level comparison (two W. acetigignens strains)', '', ''],
    ['Markers present in both strains (%)', '176 / 200 (88.0%)', 'n.a. (single genome)'],
    ['Markers present in one strain only', '5 / 200 (2.5%)', 'n.a.'],
    ['Markers not detected in either strain', '19 / 200 (9.5%)', '5 / 200 (2.5%)'],
    ['Identity between strains, median', 0.9882, 'n.a.'],
    ['', ''],
    ['Cross-species control (specificity)', '', ''],
    ['W. acetigignens markers vs D. hominis genome (median identity)', 0.272, ''],
    ['W. acetigignens markers vs D. hominis genome (>=90% id)', '0 / 200', ''],
    ['D. hominis markers vs both W. acetigignens genomes (median identity)', '0.274 / 0.274', ''],
    ['D. hominis markers vs both W. acetigignens genomes (>=90% id)', '0 / 200 / 0 / 200', ''],
]
for r in rows:
    ws2.append(r)

# ---- 格式化 ----
hfill = PatternFill("solid", fgColor="DDEBF7")
for c in ws[1]:
    c.font = Font(bold=True)
    c.fill = hfill
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
ws.freeze_panes = "A2"
widths = [20, 9, 18, 16, 40, 12, 10, 9] + [11] * (6 * 3)
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
for c in ws2['A']:
    if c.value and isinstance(c.value, str) and c.value not in ('',):
        c.font = Font(bold=False)
for c in ws2[1]:
    c.font = Font(bold=True, size=12)
for c in ws2[3]:
    c.font = Font(bold=True)
for i, w in enumerate([48, 34, 30], 1):
    ws2.column_dimensions[get_column_letter(i)].width = w
for row in ws2.iter_rows():
    for c in row:
        c.alignment = Alignment(vertical='top', wrap_text=False)

wb.save(OUT)
print("已保存:", OUT)
print("明细行数:", ws.max_row - 1)
